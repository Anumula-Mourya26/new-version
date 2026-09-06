import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_doc(txt_file, docx_file, is_strategic):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54); s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.54); s.right_margin = Cm(2.54)
        
    styles = doc.styles
    title_style = styles['Title']
    title_style.font.name = 'Calibri'; title_style.font.size = Pt(36); title_style.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20); title_style.font.bold = True
    subtitle_style = styles['Subtitle']
    subtitle_style.font.name = 'Calibri'; subtitle_style.font.size = Pt(20); subtitle_style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); subtitle_style.font.italic = True
    
    h1_style = styles.add_style('AnnSetu_H1', WD_STYLE_TYPE.PARAGRAPH)
    h1_style.base_style = styles['Heading 1']
    h1_style.font.name = 'Calibri'; h1_style.font.size = Pt(18); h1_style.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20); h1_style.font.bold = True
    
    h2_style = styles.add_style('AnnSetu_H2', WD_STYLE_TYPE.PARAGRAPH)
    h2_style.base_style = styles['Heading 2']
    h2_style.font.name = 'Calibri'; h2_style.font.size = Pt(14); h2_style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); h2_style.font.bold = True
    
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'; normal_style.font.size = Pt(11); normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_paragraph('\n' * 5)
    
    if is_strategic:
        title = doc.add_paragraph('SIH26032 — STRATEGIC ANALYSIS & JURY EVALUATION', style='Title')
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = doc.add_paragraph('Competitive Positioning, Concept Scoring, and the Winning Concept', style='Subtitle')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        desc = doc.add_paragraph("Ministry of Consumer Affairs, Food & Public Distribution\nTrack: Software · Theme: Smart Automation\nSeptember 2026\n")
        desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        title = doc.add_paragraph('ANNSETU', style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = doc.add_paragraph('"Bridge to the Grain"', style='Subtitle'); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        desc = doc.add_paragraph("Engineering Blueprint — Live Gate & Queue Transparency Layer for MSP Procurement\nSIH26032 · Ministry of Consumer Affairs, Food & Public Distribution\nTrack: Software · Theme: Smart Automation\nPrepared as a Principal-Engineer build specification — not generic hackathon advice.\nSeptember 2026\n")
        desc.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    text = open(txt_file, encoding='utf-8').read()
    if is_strategic:
        text = re.sub(r'^.*?Diagnostic Analysis', 'Diagnostic Analysis', text, flags=re.DOTALL)
    else:
        text = re.sub(r'^.*?1\. Product', '1. Product', text, flags=re.DOTALL)
    
    for line in text.split('\n'):
        line = line.strip()
        if not line: continue
        if line.startswith('\x0c'):
            doc.add_page_break()
            line = line.lstrip('\x0c').strip()
            if not line: continue
            
        bp_h1_pattern = r'^((?:[1-9]|1[0-9]|20)\. (?:Product|Complete User Workflow|System Architecture|Technology Selection|Database|API Design|Queue Engine|Prediction Engine|Notification Engine|Farmer Experience|Admin Dashboard|Codebase|Implementation Order|Code Generation Rule|Demo|Metrics|Testing|Deployment|Government Deployment Roadmap|Final Build Checklist))$'
        if (not is_strategic and re.match(bp_h1_pattern, line)) or (is_strategic and (line == 'Diagnostic Analysis' or line.startswith('10 Candidate') or line.startswith('Competitor Attack') or line.startswith('WINNING CONCEPT'))):
            doc.add_paragraph(''.join(c for c in line if ord(c) >= 32), style='AnnSetu_H1')
        elif (line.endswith(':') and len(line) < 60) or (is_strategic and re.match(r'^\d+\.', line)) or (not is_strategic and re.match(r'^[A-Z]\.', line)):
            doc.add_paragraph(''.join(c for c in line if ord(c) >= 32), style='AnnSetu_H2')
        elif line.startswith('•'):
            doc.add_paragraph(''.join(c for c in line[1:].strip() if ord(c) >= 32), style='List Bullet')
        else:
            doc.add_paragraph(''.join(c for c in line if ord(c) >= 32))

    doc.save(docx_file)

create_doc('blueprint_full.txt', 'AnnSetu_Engineering_Blueprint.docx', False)
create_doc('sih_full.txt', 'SIH26032_Strategic_Analysis.docx', True)
